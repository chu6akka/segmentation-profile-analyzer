import pandas as pd
import pytest
from openpyxl import load_workbook

from comparison import compare_profiles
from export import export_csv, export_xlsx, export_comparison_xlsx
from features import analyze_author
from parser import parse_chat


def test_exports_and_comparison(tmp_path):
    path = tmp_path / "chat.txt"
    path.write_text('[11.06.2026 14:12] A: =1+1', encoding="utf-8")
    result = analyze_author(parse_chat(str(path)), "A")
    export_csv(result, tmp_path / "profile.csv")
    assert pd.read_csv(tmp_path / "profile.csv").iloc[0].author == "A"
    export_xlsx(result, tmp_path / "profile.xlsx")
    book = load_workbook(tmp_path / "profile.xlsx")
    assert book.sheetnames == ["Profile", "Messages", "Series"]
    assert book["Messages"]["D2"].value == "=1+1"
    assert book["Messages"]["D2"].data_type == "s"
    comparison = compare_profiles(result, result)
    assert comparison.absolute_difference.dropna().eq(0).all()
    export_comparison_xlsx(result, result, tmp_path / "comparison.xlsx")
    assert load_workbook(tmp_path / "comparison.xlsx").sheetnames == ["Profile_A", "Profile_B", "Comparison"]


def test_failed_export_preserves_existing_file(tmp_path):
    path = tmp_path / "chat.txt"
    path.write_text('[11.06.2026 14:12] A: ' + 'а' * 32768, encoding="utf-8")
    result = analyze_author(parse_chat(str(path)), "A")
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"existing content")
    with pytest.raises(ValueError, match="32767"):
        export_xlsx(result, output)
    assert output.read_bytes() == b"existing content"
