import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
from openpyxl import load_workbook

from ui.main_window import MainWindow
from ui.theme import configure_app


@pytest.fixture(scope="module")
def app():
    application = QApplication.instance() or QApplication([])
    configure_app(application)
    return application


@pytest.fixture
def window(app, tmp_path):
    path = tmp_path / "chat.txt"
    path.write_text("\n".join([
        "[11.06.2026 9:12] A: раз два три четыре пять шесть семь восемь девять десять",
        "[11.06.2026 9:12] A: два слова",
        "[11.06.2026 9:13] B: Да.",
        "[11.06.2026 9:14] A: привет",
    ]), encoding="utf-8")
    widget = MainWindow()
    widget.picker.load_path(str(path))
    widget.show()
    app.processEvents()
    yield widget, path
    widget.close()


def test_analysis_sorting_charts_and_invalidation(window, app):
    widget, _ = window
    QTest.mouseClick(widget.analyze_button, Qt.MouseButton.LeftButton)
    app.processEvents()
    assert widget.analysis.profile["message_count"] == 3
    assert len(widget.charts.figure.axes) == 3
    widget.messages.sortByColumn(4, Qt.SortOrder.AscendingOrder)
    model = widget.messages.model()
    assert [model.data(model.index(i, 4)) for i in range(3)] == ["1", "2", "10"]
    widget.picker.authors.setCurrentText("B")
    assert widget.analysis is None
    assert not widget.xlsx_button.isEnabled()
    assert widget.messages.model().rowCount() == 0


def test_gui_exports_and_comparison(window, app, tmp_path, monkeypatch):
    widget, source = window
    widget.analyze()
    target = tmp_path / "out.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **kw: (str(target), ""))
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **kw: None)
    widget.save(True)
    assert load_workbook(target).sheetnames == ["Profile", "Messages", "Series"]
    tab = widget.comparison
    tab.a.load_path(str(source))
    tab.b.load_path(str(source))
    tab.b.authors.setCurrentText("B")
    tab.analyze()
    assert tab.results[0].profile["message_count"] == 3
    assert tab.results[1].profile["message_count"] == 1
    assert tab.table.model().rowCount() == 16
    tab.save(True)
    assert load_workbook(target).sheetnames == ["Author_A", "Author_B", "Between_Author_Comparison"]
    tab.b.authors.setCurrentText("A")
    assert tab.results is None and not tab.xlsx_button.isEnabled()

